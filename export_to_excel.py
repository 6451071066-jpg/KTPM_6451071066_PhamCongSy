import os
import sys
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Đảm bảo mã hóa UTF-8 an toàn trên mọi môi trường
if sys.stdout and hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")


def export_test_cases_to_excel(filename="BaoCao_Test_Login_UTC.xlsx"):
    """Xuất báo cáo kết quả kiểm thử tự động toàn diện ra file Excel với 2 Sheets chuyên nghiệp."""
    wb = openpyxl.Workbook()

    # Định nghĩa bảng màu thiết kế chuyên nghiệp (Corporate Theme)
    COLOR_PRIMARY = "1F4E78"       # Xanh navy đậm cho Tiêu đề và Header
    COLOR_SECONDARY = "2F5597"     # Xanh trung tính cho Sub-header
    COLOR_ACCENT = "D9E1F2"        # Xanh nhạt highlight
    COLOR_ZEBRA = "F9FAFC"         # Màu nền hàng chẵn
    COLOR_WHITE = "FFFFFF"         # Màu trắng
    COLOR_BORDER = "D3D3D3"        # Viền xám mỏng
    
    # Màu trạng thái
    FILL_PASSED = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")  # Xanh lá nhạt
    FONT_PASSED = Font(name="Segoe UI", size=10, bold=True, color="385723")

    FILL_FAILED = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")  # Đỏ cam nhạt
    FONT_FAILED = Font(name="Segoe UI", size=10, bold=True, color="C00000")

    FILL_SKIPPED = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid") # Vàng nhạt
    FONT_SKIPPED = Font(name="Segoe UI", size=10, bold=True, color="7F6000")

    thin_border = Border(
        left=Side(style="thin", color=COLOR_BORDER),
        right=Side(style="thin", color=COLOR_BORDER),
        top=Side(style="thin", color=COLOR_BORDER),
        bottom=Side(style="thin", color=COLOR_BORDER),
    )

    # =========================================================================
    # SHEET 1: TỔNG QUAN & THỐNG KÊ (DASHBOARD & METRICS)
    # =========================================================================
    ws1 = wb.active
    ws1.title = "Tổng Quan & Thống Kê"
    ws1.views.sheetView[0].showGridLines = True

    # 1.1 Banner Tiêu đề
    ws1.merge_cells("A1:G1")
    title_cell = ws1["A1"]
    title_cell.value = "BÁO CÁO THỰC THI KIỂM THỬ TỰ ĐỘNG - HỆ THỐNG VĂN PHÒNG ĐIỆN TỬ UTC"
    title_cell.font = Font(name="Segoe UI", size=15, bold=True, color=COLOR_WHITE)
    title_cell.fill = PatternFill(start_color=COLOR_PRIMARY, end_color=COLOR_PRIMARY, fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 40

    ws1.merge_cells("A2:G2")
    subtitle_cell = ws1["A2"]
    subtitle_cell.value = "DỰ ÁN AUTOMATION TESTING E2E (PAGE OBJECT MODEL - SELENIUM & PYTEST)"
    subtitle_cell.font = Font(name="Segoe UI", size=10, italic=True, color="D9E1F2")
    subtitle_cell.fill = PatternFill(start_color=COLOR_SECONDARY, end_color=COLOR_SECONDARY, fill_type="solid")
    subtitle_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[2].height = 24

    # 1.2 Headers phân vùng
    ws1.merge_cells("A3:C3")
    sec1 = ws1["A3"]
    sec1.value = "THÔNG TIN DỰ ÁN & MÔI TRƯỜNG KIỂM THỬ"
    sec1.font = Font(name="Segoe UI", size=11, bold=True, color=COLOR_WHITE)
    sec1.fill = PatternFill(start_color=COLOR_PRIMARY, end_color=COLOR_PRIMARY, fill_type="solid")
    sec1.alignment = Alignment(horizontal="center", vertical="center")

    ws1.merge_cells("E3:G3")
    sec2 = ws1["E3"]
    sec2.value = "KẾT QUẢ THỰC THI THỰC TẾ"
    sec2.font = Font(name="Segoe UI", size=11, bold=True, color=COLOR_WHITE)
    sec2.fill = PatternFill(start_color=COLOR_PRIMARY, end_color=COLOR_PRIMARY, fill_type="solid")
    sec2.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[3].height = 26

    # 1.3 Dữ liệu thông tin dự án
    now_str = datetime.datetime.now().strftime("%d/%m/%Y %H:%M:%S")
    project_info = [
        ("Hệ thống kiểm thử (AUT):", "Văn phòng điện tử UTC (https://vanphongdientu.utc.edu.vn/Login)"),
        ("Môi trường kiểm thử:", "Google Chrome (Latest) / Windows 10/11"),
        ("Framework & Ngôn ngữ:", "Selenium WebDriver 4.20+ / Python 3.13 / Pytest 9.1+"),
        ("Người thực hiện kiểm thử:", "Phạm Công Sỹ - MSSV: 6451071066"),
        ("Mã dự án / Repository:", "KTPM_6451071066_PhamCongSy"),
        ("Báo cáo đồ họa & Log:", "Allure Framework 2.16+ / Microsoft Excel (OpenPyXL)"),
        ("Thời gian thực thi:", now_str),
    ]

    for idx, (label, val) in enumerate(project_info, start=4):
        c1 = ws1.cell(row=idx, column=1, value=label)
        c1.font = Font(name="Segoe UI", size=10, bold=True, color="262626")
        c1.fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
        c1.border = thin_border
        
        ws1.merge_cells(start_row=idx, start_column=2, end_row=idx, end_column=3)
        c2 = ws1.cell(row=idx, column=2, value=val)
        c2.font = Font(name="Segoe UI", size=10)
        c2.border = thin_border
        ws1.cell(row=idx, column=3).border = thin_border
        ws1.row_dimensions[idx].height = 22

    # 1.4 Bảng thống kê kết quả kiểm thử
    metrics = [
        ("Chỉ Số Đo Lường", "Số Lượng", "Tỷ Lệ (%)"),
        ("Tổng số Test Cases", 16, "100.0%"),
        ("Passed (Đạt yêu cầu)", 14, "87.5%"),
        ("Failed (Phát hiện lỗi / Bug)", 1, "6.25%"),
        ("Skipped (Chờ tài khoản thật)", 1, "6.25%"),
    ]

    for idx, row_data in enumerate(metrics, start=4):
        ws1.row_dimensions[idx].height = 22
        for col_idx, val in enumerate(row_data, start=5):
            cell = ws1.cell(row=idx, column=col_idx, value=val)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if idx == 4:
                cell.font = Font(name="Segoe UI", size=10, bold=True, color=COLOR_WHITE)
                cell.fill = PatternFill(start_color=COLOR_SECONDARY, end_color=COLOR_SECONDARY, fill_type="solid")
            else:
                if col_idx == 5:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                    cell.font = Font(name="Segoe UI", size=10, bold=True)
                elif "Passed" in row_data[0]:
                    cell.fill = FILL_PASSED
                    cell.font = FONT_PASSED
                elif "Failed" in row_data[0]:
                    cell.fill = FILL_FAILED
                    cell.font = FONT_FAILED
                elif "Skipped" in row_data[0]:
                    cell.fill = FILL_SKIPPED
                    cell.font = FONT_SKIPPED
                else:
                    cell.font = Font(name="Segoe UI", size=10, bold=True)

    # 1.5 Kết luận đánh giá
    ws1.merge_cells("A12:G12")
    c_conc = ws1["A12"]
    c_conc.value = (
        "👉 KẾT LUẬN: Bộ kiểm thử tự động gồm 16 kịch bản bao phủ trọn vẹn màn hình Đăng nhập UTC. "
        "Phát hiện 01 Defect thực tế tại TC11 (không trim khoảng trắng), bảo vệ thành công trước tấn công SQL Injection (TC14) "
        "và XSS (TC16), hỗ trợ submit bằng phím Enter (TC15), cùng cơ chế Anti-Brute Force Captcha hoạt động chính xác (TC10)."
    )
    c_conc.font = Font(name="Segoe UI", size=10, italic=True, bold=True, color="1F4E78")
    c_conc.fill = PatternFill(start_color="EDF2F8", end_color="EDF2F8", fill_type="solid")
    c_conc.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    c_conc.border = thin_border
    ws1.row_dimensions[12].height = 42

    # Điều chỉnh độ rộng cột Sheet 1
    col_widths_s1 = {"A": 26, "B": 24, "C": 26, "D": 4, "E": 28, "F": 12, "G": 14}
    for col_letter, width in col_widths_s1.items():
        ws1.column_dimensions[col_letter].width = width

    # =========================================================================
    # SHEET 2: CHI TIẾT TEST CASES (DETAILED TEST SPECIFICATIONS & RESULTS)
    # =========================================================================
    ws2 = wb.create_sheet(title="Chi Tiết Test Cases")
    ws2.views.sheetView[0].showGridLines = True

    # 2.1 Tiêu đề Sheet 2
    ws2.merge_cells("A1:K1")
    title2 = ws2["A1"]
    title2.value = "BẢNG ĐẶC TẢ VÀ KẾT QUẢ THỰC THI 16 KỊCH BẢN KIỂM THỬ (TEST CASES SPECIFICATION)"
    title2.font = Font(name="Segoe UI", size=14, bold=True, color=COLOR_WHITE)
    title2.fill = PatternFill(start_color=COLOR_PRIMARY, end_color=COLOR_PRIMARY, fill_type="solid")
    title2.alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 36

    # 2.2 Headers cho bảng
    headers = [
        "Mã TC",
        "Phân Hệ / Tính Năng",
        "Mô Tả Kịch Bản",
        "Độ Nghiêm Trọng",
        "Điều Kiện Tiền Đề",
        "Các Bước Thực Hiện",
        "Dữ Liệu Thử Nghiệm (Test Data)",
        "Kết Quả Mong Đợi",
        "Kết Quả Thực Tế",
        "Trạng Thái",
        "Ghi Chú / Phân Loại Lỗi"
    ]

    ws2.row_dimensions[2].height = 28
    for col_idx, header_text in enumerate(headers, start=1):
        cell = ws2.cell(row=2, column=col_idx, value=header_text)
        cell.font = Font(name="Segoe UI", size=10, bold=True, color=COLOR_WHITE)
        cell.fill = PatternFill(start_color=COLOR_SECONDARY, end_color=COLOR_SECONDARY, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    # 2.3 Danh sách dữ liệu chi tiết của 16 Test Cases
    test_cases_data = [
        [
            "TC01",
            "Bảo mật Giao diện",
            "Kiểm tra ô mật khẩu ẩn ký tự (type='password')",
            "Critical",
            "Trang /Login tải hoàn tất",
            "1. Mở trang Login UTC\n2. Lấy thẻ input mật khẩu\n3. Kiểm tra thuộc tính type",
            "N/A",
            "Trường mật khẩu có thuộc tính type='password', ký tự nhập vào hiển thị dạng dấu chấm/sao",
            "Thuộc tính type của trường password là 'password', ký tự được ẩn an toàn",
            "PASSED",
            "Đạt tiêu chuẩn an toàn bảo mật hiển thị"
        ],
        [
            "TC02",
            "Tương tác Giao diện",
            "Kiểm tra bật/tắt checkbox 'Giữ tôi luôn đăng nhập'",
            "Minor",
            "Checkbox hiển thị trên form đăng nhập",
            "1. Đọc trạng thái ban đầu\n2. Click vào nhãn checkbox\n3. Kiểm tra trạng thái is_selected()",
            "Thao tác click chuột",
            "Trạng thái checkbox chuyển đổi thành công (Checked <-> Unchecked)",
            "Checkbox đảo trạng thái chính xác sau thao tác click",
            "PASSED",
            "Tương tác UI ổn định"
        ],
        [
            "TC03",
            "Validation Form",
            "Bấm Đăng nhập khi để trống cả Tên đăng nhập và Mật khẩu",
            "Normal",
            "Trang /Login hiển thị 2 trường trống",
            "1. Không điền dữ liệu vào form\n2. Nhấn nút 'Đăng nhập'\n3. Kiểm tra thông báo lỗi",
            "username=''\nuserpwd=''",
            "Hệ thống chặn submit và báo lỗi: 'Bạn chưa nhập tên đăng nhập'",
            "Hiển thị thông báo: 'Bạn chưa nhập tên đăng nhập'",
            "PASSED",
            "Validation form hoạt động đúng"
        ],
        [
            "TC04",
            "Validation Form",
            "Nhập Tên đăng nhập nhưng để trống Mật khẩu",
            "Normal",
            "Trang /Login sẵn sàng",
            "1. Điền Tên đăng nhập hợp lệ\n2. Để trống ô Mật khẩu\n3. Nhấn nút 'Đăng nhập'",
            "username='sinhvien_utc'\nuserpwd=''",
            "Hệ thống chặn submit và báo lỗi: 'Bạn chưa nhập mật khẩu'",
            "Hiển thị thông báo: 'Bạn chưa nhập mật khẩu'",
            "PASSED",
            "Validation form hoạt động đúng"
        ],
        [
            "TC05",
            "Validation Form",
            "Để trống Tên đăng nhập nhưng có nhập Mật khẩu",
            "Normal",
            "Trang /Login sẵn sàng",
            "1. Để trống ô Tên đăng nhập\n2. Điền ô Mật khẩu\n3. Nhấn nút 'Đăng nhập'",
            "username=''\nuserpwd='MatKhau123@'",
            "Hệ thống chặn submit và báo lỗi: 'Bạn chưa nhập tên đăng nhập'",
            "Hiển thị thông báo: 'Bạn chưa nhập tên đăng nhập'",
            "PASSED",
            "Validation form hoạt động đúng"
        ],
        [
            "TC06",
            "Xác thực (Authentication)",
            "Đăng nhập với tài khoản hoặc mật khẩu không chính xác",
            "Critical",
            "Trang /Login sẵn sàng",
            "1. Nhập tài khoản không tồn tại / sai mật khẩu\n2. Nhấn nút 'Đăng nhập'\n3. Kiểm tra phản hồi",
            "username='sinhvien_utc_invalid'\nuserpwd='MatKhauSai@999'",
            "Hệ thống từ chối đăng nhập và báo lỗi: 'Tài khoản hoặc mật khẩu không đúng.'",
            "Hiển thị thông báo: 'Tài khoản hoặc mật khẩu không đúng.'",
            "PASSED",
            "Cơ chế xác thực từ chối thông tin sai an toàn"
        ],
        [
            "TC07",
            "Điều hướng (Navigation)",
            "Kiểm tra liên kết 'Bạn quên mật khẩu đăng nhập ?'",
            "Normal",
            "Trang /Login hiển thị liên kết quên mật khẩu",
            "1. Click vào liên kết quên mật khẩu\n2. Chờ chuyển trang\n3. Kiểm tra URL đích",
            "Thao tác click link /Login/GetPass",
            "Trình duyệt điều hướng chính xác sang trang cấp lại mật khẩu (/Login/GetPass)",
            "Chuyển hướng thành công tới URL chứa '/Login/GetPass'",
            "PASSED",
            "Liên kết điều hướng hoạt động tốt"
        ],
        [
            "TC08",
            "Cổng xác thực SSO",
            "Kiểm tra điều hướng nút 'Đăng nhập bằng e-mail UTC'",
            "Normal",
            "Nút đăng nhập SSO hiển thị",
            "1. Click nút 'Đăng nhập bằng e-mail UTC'\n2. Chờ chuyển hướng OAuth\n3. Kiểm tra URL đích",
            "Thao tác click nút SSO UTC",
            "Hệ thống chuyển hướng sang cổng xác thực Google OAuth (accounts.google.com)",
            "Trình duyệt chuyển hướng chính xác đến accounts.google.com",
            "PASSED",
            "Tích hợp Google OAuth hoạt động chuẩn"
        ],
        [
            "TC09",
            "Luồng chính (Happy Path)",
            "Đăng nhập thành công với tài khoản UTC hợp lệ",
            "Blocker",
            "Có tài khoản và mật khẩu thật của sinh viên/giảng viên UTC",
            "1. Nhập username thật\n2. Nhập password thật\n3. Bấm Đăng nhập\n4. Kiểm tra URL chuyển vào Dashboard",
            "Tài khoản UTC thực tế",
            "Đăng nhập thành công, URL chuyển hướng thoát khỏi /Login vào trang nội bộ văn phòng",
            "Skipped: Cần cung cấp tài khoản và mật khẩu UTC thực tế để chạy case này",
            "SKIPPED",
            "Được đánh dấu @pytest.mark.skip, bảo vệ thông tin mật cá nhân"
        ],
        [
            "TC10",
            "Bảo mật chống Brute-force",
            "Kích hoạt mã bảo mật (Captcha) khi nhập sai liên tiếp từ 3 lần",
            "Critical",
            "Trang /Login sẵn sàng",
            "1. Thực hiện 3 lần đăng nhập liên tiếp với thông tin sai\n2. Kiểm tra giao diện form đăng nhập",
            "username='sinhvien_sai_tk'\nuserpwd='MatKhauSai@0..2'",
            "Hệ thống tự động kích hoạt hiển thị ảnh Captcha và ô nhập Captcha để chống Brute-Force",
            "Ảnh Captcha và ô nhập Captcha xuất hiện chính xác sau lần sai thứ 3",
            "PASSED",
            "Cơ chế chống dò quét mật khẩu (Anti-Brute Force) hoạt động rất tốt"
        ],
        [
            "TC11",
            "Validation Form (BUG)",
            "Nhập toàn ký tự khoảng trắng (Space) vào Tên đăng nhập và Mật khẩu",
            "Normal",
            "Trang /Login sẵn sàng",
            "1. Nhập chuỗi chỉ gồm khoảng trắng vào 2 trường\n2. Nhấn nút Đăng nhập\n3. Kiểm tra thông báo",
            "username='   '\nuserpwd='   '",
            "Hệ thống tự động trim khoảng trắng và báo lỗi: 'Bạn chưa nhập tên đăng nhập'",
            "Hệ thống không trim khoảng trắng, gửi dữ liệu lên server và báo: 'Tài khoản hoặc mật khẩu không đúng.'",
            "FAILED",
            "[DEFECT]: Phát hiện Bug thực tế - Hệ thống thiếu xử lý trim whitespace ở client/server"
        ],
        [
            "TC12",
            "Giá trị biên (BVA)",
            "Kiểm tra nhập dữ liệu độ dài tối thiểu (1 ký tự)",
            "Minor",
            "Trang /Login sẵn sàng",
            "1. Nhập username 1 ký tự\n2. Nhập password 1 ký tự\n3. Nhấn nút Đăng nhập",
            "username='a'\nuserpwd='1'",
            "Hệ thống xử lý an toàn giá trị biên dưới, từ chối và báo lỗi: 'Tài khoản hoặc mật khẩu không đúng.'",
            "Hiển thị thông báo: 'Tài khoản hoặc mật khẩu không đúng.'",
            "PASSED",
            "Xử lý biên dưới an toàn, không có lỗi giao diện"
        ],
        [
            "TC13",
            "Kiểm thử chịu tải độ dài",
            "Kiểm tra nhập chuỗi cực dài (500 ký tự) vào các ô nhập liệu",
            "Normal",
            "Trang /Login sẵn sàng",
            "1. Nhập 500 ký tự vào username\n2. Nhập 500 ký tự vào password\n3. Nhấn nút Đăng nhập",
            "username='user'*125 (500 chars)\nuserpwd='pass'*125 (500 chars)",
            "Hệ thống không bị treo hoặc crash lỗi máy chủ (Error 500), phản hồi lỗi an toàn",
            "Hiển thị thông báo lỗi 'Tài khoản hoặc mật khẩu không đúng.', hệ thống ổn định",
            "PASSED",
            "Hệ thống có bộ đệm tốt, không xảy ra buffer overflow / HTTP 500"
        ],
        [
            "TC14",
            "Bảo mật chống Injection",
            "Kiểm tra phòng chống tấn công SQL Injection vào form đăng nhập",
            "Critical",
            "Trang /Login sẵn sàng",
            "1. Nhập chuỗi payload SQL Injection (' OR '1'='1)\n2. Nhấn Đăng nhập\n3. Kiểm tra phản hồi bảo mật",
            "username=\"' OR '1'='1\"\nuserpwd=\"' OR '1'='1\"",
            "Hệ thống ngăn chặn tấn công SQLi, không rò rỉ cú pháp cơ sở dữ liệu và từ chối truy cập an toàn",
            "Hệ thống bảo vệ an toàn, không bypass xác thực, báo lỗi: 'Tài khoản hoặc mật khẩu không đúng.'",
            "PASSED",
            "Cơ chế phòng chống SQL Injection đạt chuẩn an toàn thông tin"
        ],
        [
            "TC15",
            "Trải nghiệm người dùng (UX)",
            "Kiểm tra gửi form đăng nhập bằng phím ENTER thay vì click nút Đăng nhập",
            "Normal",
            "Trang /Login sẵn sàng, con trỏ tại trường nhập liệu",
            "1. Nhập Tên đăng nhập và Mật khẩu\n2. Nhấn phím ENTER trên bàn phím tại ô Mật khẩu\n3. Kiểm tra form submit",
            "username='sinhvien_utc_enter'\nuserpwd='MatKhauSai@123'\nAction: Keys.ENTER",
            "Form tự động kích hoạt gửi request đăng nhập khi người dùng nhấn phím ENTER",
            "Form được submit chính xác bằng phím ENTER, nhận thông báo phản hồi từ hệ thống",
            "PASSED",
            "Hỗ trợ tiện ích điều hướng bàn phím (Keyboard Usability & Accessibility)"
        ],
        [
            "TC16",
            "Bảo mật chống Injection",
            "Kiểm tra phòng chống tấn công Cross-Site Scripting (XSS) trong ô nhập liệu",
            "Critical",
            "Trang /Login sẵn sàng",
            "1. Nhập mã script XSS (<script>alert('XSS_ATTACK')</script>)\n2. Nhấn nút Đăng nhập\n3. Kiểm tra mã độc",
            "username=\"<script>alert('XSS_ATTACK')</script>\"\nuserpwd=\"Password123@\"",
            "Hệ thống làm sạch (Sanitize) mã độc, không thực thi popup JavaScript và trả về thông báo an toàn",
            "Không có popup script thực thi, hệ thống xử lý chuỗi an toàn và báo lỗi đăng nhập",
            "PASSED",
            "Phòng chống tấn công XSS thành công"
        ],
    ]

    for row_idx, row_values in enumerate(test_cases_data, start=3):
        # Thiết lập màu nền xen kẽ (Zebra Striping)
        is_even = (row_idx % 2 == 0)
        default_fill = PatternFill(start_color=COLOR_ZEBRA, end_color=COLOR_ZEBRA, fill_type="solid") if is_even else PatternFill(start_color=COLOR_WHITE, end_color=COLOR_WHITE, fill_type="solid")

        ws2.row_dimensions[row_idx].height = 40

        for col_idx, val in enumerate(row_values, start=1):
            cell = ws2.cell(row=row_idx, column=col_idx, value=val)
            cell.font = Font(name="Segoe UI", size=9)
            cell.border = thin_border
            cell.fill = default_fill

            # Căn chỉnh văn bản theo từng loại cột
            if col_idx in [1, 4]:  # Mã TC, Độ nghiêm trọng
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in [2, 10]:  # Phân hệ, Trạng thái
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

            # Định dạng riêng cho Mã TC
            if col_idx == 1:
                cell.font = Font(name="Segoe UI", size=10, bold=True, color="1F4E78")

            # Định dạng riêng cho cột Độ nghiêm trọng
            if col_idx == 4:
                if val == "Blocker":
                    cell.font = Font(name="Segoe UI", size=9, bold=True, color="7F00FF")
                elif val == "Critical":
                    cell.font = Font(name="Segoe UI", size=9, bold=True, color="C00000")
                elif val == "Normal":
                    cell.font = Font(name="Segoe UI", size=9, color="262626")
                elif val == "Minor":
                    cell.font = Font(name="Segoe UI", size=9, color="595959")

            # Định dạng riêng cho cột Trạng Thái (Status)
            if col_idx == 10:
                if val == "PASSED":
                    cell.fill = FILL_PASSED
                    cell.font = FONT_PASSED
                elif val == "FAILED":
                    cell.fill = FILL_FAILED
                    cell.font = FONT_FAILED
                elif val == "SKIPPED":
                    cell.fill = FILL_SKIPPED
                    cell.font = FONT_SKIPPED

            # Tô đậm ghi chú lỗi Bug
            if col_idx == 11 and "[DEFECT]" in str(val):
                cell.font = Font(name="Segoe UI", size=9, bold=True, color="C00000")

    # 2.4 Căn chỉnh độ rộng các cột của Sheet 2
    col_widths_s2 = {
        "A": 9,   # Mã TC
        "B": 22,  # Phân hệ
        "C": 30,  # Mô tả kịch bản
        "D": 15,  # Độ nghiêm trọng
        "E": 26,  # Điều kiện tiền đề
        "F": 34,  # Các bước thực hiện
        "G": 30,  # Dữ liệu thử nghiệm
        "H": 34,  # Kết quả mong đợi
        "I": 34,  # Kết quả thực tế
        "J": 14,  # Trạng thái
        "K": 32   # Ghi chú
    }

    for col_letter, width in col_widths_s2.items():
        ws2.column_dimensions[col_letter].width = width

    # Lưu workbook
    wb.save(filename)
    print(f"Xuat bao cao kiem thu thanh cong ra file: {filename}")


if __name__ == "__main__":
    export_test_cases_to_excel()
